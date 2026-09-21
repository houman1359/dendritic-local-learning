"""Lossless CSV presentation with separately labeled scientific procedures.

No statistic or model choice is recomputed. Optional references require an
explicit hash-bound manifest and an explicit prior/supervision scope per table.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEETS = {
    "curves": "Risk curves",
    "teacher_curves": "Teacher risks",
    "paired_ratios": "Architecture ratios",
    "teacher_pairs": "Teacher architecture ratios",
    "primary_comparisons": "Two primary comparisons",
    "sample_size_ratios": "Sample size ratios",
    "teacher_sample_size_ratios": "Teacher sample size ratios",
    "budget_changes": "Budget changes",
    "forecast_errors": "Numerical forecasts",
    "architecture_predictions": "Architecture forecasts",
    "all_endpoints": "All fitted endpoints",
}
SCOPES = {
    "generic": "Generic TRAIN-only fitting; supplied block/rank/input-law priors, no known radial profile or private projectors.",
    "known_profile": "Stronger known-profile learner: scalar TRAIN-estimated planes plus analytic radial profile, normalization and input law; no private projector.",
    "oracle": "Private-plane approximation witness; private projector information is supplied. This is not a learned result.",
    "polynomial": "Ordinary polynomial least squares using scalar TRAIN labels, declared polynomial degree and supplied blocks; degree prior is explicit.",
    "mixed_reference": "Mixed reference table: each row must declare one of known_profile/oracle/polynomial in its declared scope column; priors differ and no rows are relabeled generic.",
}
INTEGER = re.compile(r"[+-]?\d+\Z")
NUMBER = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z")
BORDER = Border(
    **{
        k: Side(style="thin", color="D8DCE0")
        for k in ("left", "right", "top", "bottom")
    }
)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    return json.loads(Path(path).read_text())


def parse(value):
    if value == "":
        return None
    if value in ("True", "False"):
        return value == "True"
    if INTEGER.fullmatch(value):
        return int(value) if len(value.lstrip("+-")) <= 15 else value
    if NUMBER.fullmatch(value):
        number = float(value)
        return number if math.isfinite(number) else value
    return value


def csv_table(path):
    with Path(path).open(newline="") as f:
        all_rows = list(csv.reader(f))
    if not all_rows:
        return [], []
    headers = all_rows[0]
    if len(headers) != len(set(headers)) or any(
        len(r) != len(headers) for r in all_rows[1:]
    ):
        raise ValueError("CSV headers/row widths must be unambiguous")
    return headers, [[parse(v) for v in row] for row in all_rows[1:]]


def worksheet(book, title, headers, rows, *, narrative=False):
    if len(title) > 31 or title in book.sheetnames:
        raise ValueError("Unique short worksheet name required")
    sheet = book.create_sheet(title)
    sheet.append(headers or ["No declared rows"])
    for row in rows:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError("CSV text exceeds Excel cell limit")
                cell.data_type = "s"
            cell.border = BORDER
            cell.font = Font(
                name="Calibri", size=10, bold=cell.row == 1, color="222222"
            )
            cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
            cell.alignment = Alignment(
                vertical="top", wrap_text=narrative or cell.row == 1
            )
            if cell.row > 1 and isinstance(cell.value, float):
                field = headers[cell.column - 1]
                cell.number_format = (
                    "0.0000"
                    if "log" in field
                    else (
                        "0.000E+00"
                        if (
                            "mse" in field
                            or "risk" in field
                            or "objective" in field
                            or (cell.value != 0 and abs(cell.value) < 0.001)
                        )
                        else "0.0000"
                    )
                )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 32
    for i, header in enumerate(headers or ["No declared rows"], 1):
        sheet.column_dimensions[get_column_letter(i)].width = (
            90
            if header in ("detail", "path", "scope")
            else (
                68
                if header == "sha256"
                else (
                    28
                    if header in ("task", "geometry", "field", "item", "family")
                    else min(29, max(16, len(header) + 2))
                )
            )
        )
    if narrative:
        for i, row in enumerate(rows, 2):
            sheet.row_dimensions[i].height = min(
                100, 16 * (1 + max(len(str(v or "")) for v in row) // 90)
            )
    return sheet


def inputs(analysis_dir, reference_manifest=None, audit_receipts=()):
    analysis = Path(analysis_dir).resolve()
    receipt = read(analysis / "analysis_receipt.json")
    complete_paths = [
        Path(p) for p in receipt["input_bindings"] if Path(p).name == "complete.json"
    ]
    if len(complete_paths) != 1:
        raise ValueError("One bound campaign completion path required")
    root = complete_paths[0].parent
    if receipt.get("status") != "complete":
        raise ValueError("Complete sealed analysis required")
    complete = read(root / "complete.json")
    if complete.get("status") != "complete":
        raise ValueError("Complete campaign required")
    bindings = {}
    tables = []

    def bind(path, digest=None):
        path = Path(path).resolve()
        current = sha(path)
        if digest is not None and current != digest:
            raise ValueError(f"Source artifact changed: {path}")
        bindings[str(path)] = current
        return path

    bind(analysis / "analysis_receipt.json")
    bind(root / "complete.json", receipt["input_bindings"][str(root / "complete.json")])
    initialized = read(root / "initialized.json")
    bind(
        root / "initialized.json",
        receipt["input_bindings"][str(root / "initialized.json")],
    )
    for name in ("config.json", "protocol.md", "parameter_inventory.json"):
        bind(root / name, initialized["bindings"][name])
    if set(receipt["table_rows"]) != set(SHEETS):
        raise ValueError("Every declared generic analysis table must be exported")
    for name, count in receipt["table_rows"].items():
        path = bind(analysis / (name + ".csv"), receipt["outputs"][name + ".csv"])
        headers, rows = csv_table(path)
        if len(rows) != count:
            raise ValueError("CSV row count differs from analysis receipt")
        tables.append(
            {
                "name": SHEETS[name],
                "headers": headers,
                "rows": rows,
                "path": path,
                "scope": "generic",
                "detail": SCOPES["generic"],
            }
        )
    if reference_manifest:
        manifest_path = bind(reference_manifest)
        manifest = read(manifest_path)
        for i, entry in enumerate(manifest["tables"]):
            scope = entry["scope"]
            if scope not in SCOPES or scope == "generic":
                raise ValueError(
                    "Reference scope must explicitly disclose its additional prior"
                )
            path = Path(entry["path"])
            path = path if path.is_absolute() else manifest_path.parent / path
            receipt_path = Path(entry["receipt_path"])
            receipt_path = (
                receipt_path
                if receipt_path.is_absolute()
                else manifest_path.parent / receipt_path
            )
            bind(receipt_path, entry["receipt_sha256"])
            bind(path, entry["sha256"])
            reference_receipt = read(receipt_path)
            if reference_receipt.get("status") not in (
                "complete",
                "passed",
                "passed_with_documented_pointwise_roundoff_exceptions",
            ):
                raise ValueError(
                    "Reference analysis receipt must declare completed results; its separate numerical status remains explicit"
                )
            # The source receipt must itself bind the CSV, not only the supplied manifest.
            expected = reference_receipt.get("outputs", {}).get(path.name)
            if expected is None:
                expected = reference_receipt.get("input_bindings", {}).get(
                    str(path.resolve())
                )
            if expected != entry["sha256"]:
                raise ValueError("Reference CSV not bound by its declared receipt")
            headers, rows = csv_table(path)
            if scope == "mixed_reference":
                column = entry["scope_column"]
                if column not in headers:
                    raise ValueError("Mixed reference table lacks row scope column")
                values = {r[headers.index(column)] for r in rows}
                if not values <= {"known_profile", "oracle", "polynomial"}:
                    raise ValueError(
                        "Unknown or generic row in mixed privileged-reference table"
                    )
            title = entry.get("label", f"Reference {i + 1}")
            tables.append(
                {
                    "name": title,
                    "headers": headers,
                    "rows": rows,
                    "path": path.resolve(),
                    "scope": scope,
                    "detail": SCOPES[scope],
                }
            )
    audits = []
    for path in audit_receipts:
        path = bind(path)
        record = read(path)
        audits.append(
            {
                "path": str(path),
                "status": record.get("status", "Status not declared"),
                "detail": "Reported separately as supplied; no failure is relabeled as a strict pass.",
            }
        )
    bind(__file__)
    return root, receipt, tables, bindings, audits


def export(analysis_dir, output_dir, *, reference_manifest=None, audit_receipts=()):
    root, analysis, tables, bindings, audits = inputs(
        analysis_dir, reference_manifest, audit_receipts
    )
    output = Path(output_dir).resolve()
    output.mkdir(parents=True, exist_ok=False)
    book = Workbook()
    book.remove(book.active)
    book.properties.title = "Radial width: complete source-bound evidence tables"
    book.properties.creator = "Dendritic modeling project"
    is_smoke = analysis.get("scope") == "engineering_smoke_only"
    notes = [
        (
            "Scope",
            (
                "ENGINEERING SMOKE ONLY; no primary scientific inference"
                if is_smoke
                else "Complete generic radial-width evidence; every declared CSV row retained."
            ),
        ),
        (
            "Primary estimand",
            "Two teacher-mean log-risk ratios only: ReLU,N8192,P1925 width2/rank2 onquartic andwidth4/rank2 ondegree-eight. Geometric ratios are their exponentiation; arithmetic ratios are descriptive companions, never substitute primary decisions.",
        ),
        (
            "Teacher unit",
            "Eight fresh orientation teachers, two observation draws within each teacher in the full campaign. Average observations within teacher before paired resampling; observation rows are not independent teachers. Smoke configurations may use fewer.",
        ),
        (
            "Intervals",
            "10,000 shared teacher-cluster bootstrap draws; seed2026146001; linear percentiles. Pointwise95% intervals are descriptive; exactly2 primary log-ratio estimands also have nominal97.5% Bonferroni intervals. Coverage with8clusters is approximate.",
        ),
        (
            "Parameter axis",
            "Actual stored P, including the fixed soma gauge. Full/rank2/width1/width4 hit the ceiling; width3 uses one fewer scalar; widths2/5/8 use two fewer. All slopes and plots use actual P; no padding.",
        ),
        (
            "Sample size",
            "N8192/N2048 is paired within teacher/observation and uses nested TRAIN prefixes. Recipes are separately selected byN, so this compares sample size plus the declared calibration procedure, not a fixed-recipe pureN effect.",
        ),
        (
            "Risk summaries",
            "Raw arithmetic MSE and thresholded geometric MSE differ. Fixed logthreshold1e-18 and clipping counts remain inCSV. Teacher sign counts and mean effect sizes are distinct.",
        ),
        (
            "Scientific limits",
            "Favorable noiseless synthetic targets, supplied blocks, orientation-only teacher variation. No universal dendritic advantage or learned asymptotic exponent established by this workbook.",
        ),
        ("Generic fitting", SCOPES["generic"]),
        ("Known-profile learner", SCOPES["known_profile"]),
        ("Private oracle", SCOPES["oracle"]),
        ("Polynomial control", SCOPES["polynomial"]),
        (
            "Numerical audit",
            "Execution completion is distinct from independent numerical audit. All supplied audit receipts retain their own statuses; absence means not supplied. Original failures and separate reconciliations remain separate.",
        ),
        (
            "Data preservation",
            "All rows and original columns retained. Source CSVbytes remain unchanged; exported numeric/text cells are checked after workbook reopening. Display rounding changes no stored data.",
        ),
        ("Exported UTC", datetime.now(timezone.utc).isoformat()),
    ]
    worksheet(book, "Read me", ["item", "detail"], notes, narrative=True)
    status = [
        ("Analysis", analysis["status"], analysis.get("scope", "Primary analysis")),
        (
            "Main execution",
            "complete",
            "Verified global barrier; source analysis remains bound.",
        ),
        (
            "Numerical audit",
            "Not supplied" if not audits else "Separate receipts below",
            "No all-strict-pass claim inferred from model/analysis completion.",
        ),
    ]
    status.extend((a["path"], a["status"], a["detail"]) for a in audits)
    worksheet(
        book, "Execution status", ["item", "status", "detail"], status, narrative=True
    )
    inventory = read(root / "parameter_inventory.json")
    inventory_rows = []
    for entry in inventory:
        for model in entry["geometries"]:
            inventory_rows.append({"ceiling": entry["ceiling"], **model})
    ih = list(dict.fromkeys(k for row in inventory_rows for k in row))
    ir = [
        [
            (
                json.dumps(row[k], sort_keys=True)
                if isinstance(row.get(k), (list, dict))
                else row.get(k)
            )
            for k in ih
        ]
        for row in inventory_rows
    ]
    worksheet(book, "Parameter inventory", ih, ir)
    guide = []
    for table in tables:
        guide.append(
            (
                table["name"],
                table["scope"],
                len(table["rows"]),
                len(table["headers"]),
                str(table["path"]),
                table["detail"],
            )
        )
    worksheet(
        book,
        "Sheet guide",
        ["worksheet", "procedure", "rows", "columns", "path", "scope"],
        guide,
        narrative=True,
    )
    sources = [(p, d, "Verified direct input") for p, d in bindings.items()]
    sources.extend(
        (p, d, "Declared upstream sealed-analysis binding; not rehashed by exporter")
        for p, d in analysis["input_bindings"].items()
        if p not in bindings
    )
    worksheet(book, "Sources", ["path", "sha256", "verification"], sources)
    for table in tables:
        worksheet(book, table["name"], table["headers"], table["rows"])
    path = output / "radial_width_evidence.xlsx"
    with path.open("xb") as f:
        book.save(f)
    restored = load_workbook(path, read_only=True, data_only=False)
    checked = 0
    for table in tables:
        observed = list(restored[table["name"]].values)
        expected = [
            tuple(table["headers"] or ["No declared rows"]),
            *[tuple(r) for r in table["rows"]],
        ]
        if len(observed) != len(expected):
            raise ValueError("Workbook table rows changed")
        for first, second in zip(expected, observed, strict=True):
            if len(first) != len(second):
                raise ValueError("Workbook table column count changed")
            for a, b in zip(first, second, strict=True):
                if isinstance(a, float):
                    if not isinstance(b, (int, float)) or not math.isclose(
                        a, b, rel_tol=5e-15, abs_tol=0
                    ):
                        raise ValueError("Workbook numeric cell changed")
                elif a != b or (isinstance(a, bool) and not isinstance(b, bool)):
                    raise ValueError("Workbook text/bool cell changed")
                checked += 1
    if list(restored["Parameter inventory"].values) != [
        tuple(ih),
        *[tuple(row) for row in ir],
    ]:
        raise ValueError("Stored parameter inventory changed")
    restored.close()
    if any(sha(p) != d for p, d in bindings.items()):
        raise ValueError("Input changed during presentation export")
    receipt = {
        "status": "complete",
        "utc": datetime.now(timezone.utc).isoformat(),
        "source_sha256": sha(__file__),
        "analysis_scope": analysis.get("scope"),
        "workbook_path": str(path),
        "workbook_sha256": sha(path),
        "verified_input_bindings": bindings,
        "audit_receipts": audits,
        "table_rows": {t["name"]: len(t["rows"]) for t in tables},
        "table_scopes": {t["name"]: t["scope"] for t in tables},
        "roundtrip_checked_table_cells_including_headers": checked,
        "roundtrip_inventory_rows": len(ir),
        "numeric_relative_tolerance": 5e-15,
        "sheets": book.sheetnames,
        "scope": "Presentation only: no recomputed statistic, selected arm or dropped row. Generic/known-profile/oracle/polynomial scopes remain distinct; independent audit statuses are reported separately.",
    }
    with (output / "export_receipt.json").open("x") as f:
        json.dump(receipt, f, indent=2, allow_nan=False)
        f.write("\n")
    return receipt


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--analysis-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--reference-manifest", type=Path)
    parser.add_argument("--audit-receipt", type=Path, action="append", default=[])
    args = parser.parse_args()
    receipt = export(
        args.analysis_dir,
        args.output_dir,
        reference_manifest=args.reference_manifest,
        audit_receipts=args.audit_receipt,
    )
    print(
        json.dumps(
            {
                k: receipt[k]
                for k in (
                    "status",
                    "workbook_path",
                    "workbook_sha256",
                    "roundtrip_checked_table_cells_including_headers",
                )
            }
        )
    )


if __name__ == "__main__":
    main()
