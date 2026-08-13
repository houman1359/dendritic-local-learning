#!/usr/bin/env python3
"""Build styled Word documents containing each review and author response."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DRAFT_DIR = Path(__file__).resolve().parent.parent
REBUTTAL_DIR = DRAFT_DIR / "rebbutal"
REVIEWS_PATH = REBUTTAL_DIR / "rev1.txt"

REVIEWERS = ("QPha", "s2E9", "QgDY")
MARKERS = {
    "QPha": "Official Review of Submission29735 by Reviewer QPha",
    "s2E9": "Official Review of Submission29735 by Reviewer s2E9",
    "QgDY": "Official Review of Submission29735 by Reviewer QgDY",
}

DOC_PATH = REBUTTAL_DIR / "final_reviewer_reviews_and_responses.doc"
DOCX_PATH = REBUTTAL_DIR / "final_reviewer_reviews_and_responses.docx"

REVIEW_COLOR = RGBColor(0x66, 0x66, 0x66)
RESPONSE_COLOR = RGBColor(0x1F, 0x4E, 0x79)
LINE_COLOR = "9EADBA"


def clean_review(text: str) -> str:
    """Remove the OpenReview UI's trailing separator without changing the review."""
    return re.sub(r"\nAdd:\s*$", "", text.rstrip())


def load_material() -> list[tuple[str, str, str]]:
    raw = REVIEWS_PATH.read_text(encoding="utf-8")
    starts = {name: raw.index(marker) for name, marker in MARKERS.items()}
    reviews = {
        "QPha": clean_review(raw[starts["QPha"] : starts["s2E9"]]),
        "s2E9": clean_review(raw[starts["s2E9"] : starts["QgDY"]]),
        "QgDY": clean_review(raw[starts["QgDY"] :]),
    }

    material: list[tuple[str, str, str]] = []
    for reviewer in REVIEWERS:
        response_path = REBUTTAL_DIR / f"rebuttal_{reviewer}.txt"
        response = response_path.read_text(encoding="utf-8").rstrip()
        if len(response) >= 10_000:
            raise ValueError(
                f"{reviewer} response is {len(response)} characters; limit is 10,000"
            )
        forbidden = [token for token in ("**", "—", "–", "---") if token in response]
        if forbidden:
            raise ValueError(
                f"{response_path.name} still contains forbidden formatting: {forbidden}"
            )
        material.append((reviewer, reviews[reviewer], response))
    return material


def set_cell_border(paragraph, color: str = LINE_COLOR, size: str = "10") -> None:
    """Add a horizontal rule below a paragraph."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_rule(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    set_cell_border(paragraph)


def add_text_blocks(
    document: Document,
    text: str,
    *,
    size: float,
    color: RGBColor,
    spacing: float,
) -> None:
    blocks = re.split(r"\n\s*\n", text.strip())
    for block in blocks:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(spacing)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(block)
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = color


def build_docx(material: list[tuple[str, str, str]]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("NeurIPS 2026 Reviewer Responses")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.color.rgb = RESPONSE_COLOR

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Submission 29735")
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = REVIEW_COLOR

    for index, (reviewer, review, response) in enumerate(material):
        if index:
            document.add_page_break()

        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(4)
        run = heading.add_run(f"Reviewer {reviewer}")
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(15)
        run.font.color.rgb = RESPONSE_COLOR
        add_rule(document)

        label = document.add_paragraph()
        label.paragraph_format.space_after = Pt(3)
        run = label.add_run("ORIGINAL REVIEW")
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.color.rgb = REVIEW_COLOR
        add_text_blocks(
            document,
            review,
            size=8.5,
            color=REVIEW_COLOR,
            spacing=4,
        )

        add_rule(document)
        label = document.add_paragraph()
        label.paragraph_format.space_after = Pt(4)
        run = label.add_run("AUTHOR RESPONSE")
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(12)
        run.font.color.rgb = RESPONSE_COLOR
        add_text_blocks(
            document,
            response,
            size=11.5,
            color=RESPONSE_COLOR,
            spacing=7,
        )
        add_rule(document)

    document.save(DOCX_PATH)


def rtf_escape(text: str) -> str:
    """Escape Unicode text for an ASCII RTF stream."""
    pieces: list[str] = []
    for character in text:
        if character == "\\":
            pieces.append(r"\\")
        elif character == "{":
            pieces.append(r"\{")
        elif character == "}":
            pieces.append(r"\}")
        elif character == "\n":
            pieces.append(r"\line ")
        elif ord(character) < 128:
            pieces.append(character)
        else:
            encoded = character.encode("utf-16-le")
            for index in range(0, len(encoded), 2):
                unit = int.from_bytes(encoded[index : index + 2], "little")
                if unit >= 0x8000:
                    unit -= 0x10000
                pieces.append(rf"\u{unit}?")
    return "".join(pieces)


def rtf_paragraphs(
    text: str,
    *,
    font_size: int,
    color: int,
    before: int,
    after: int,
    line_spacing: int,
) -> str:
    paragraphs = re.split(r"\n\s*\n", text.strip())
    rendered: list[str] = []
    for paragraph in paragraphs:
        rendered.append(
            rf"\pard\sb{before}\sa{after}\sl{line_spacing}\slmult1"
            rf"\f0\fs{font_size}\cf{color} "
            rf"{rtf_escape(paragraph)}\par"
        )
    return "\n".join(rendered)


def build_doc(material: list[tuple[str, str, str]]) -> None:
    """Write an RTF-formatted .doc file that opens directly in Microsoft Word."""
    content = [
        r"{\rtf1\ansi\ansicpg1252\deff0\uc1",
        r"{\fonttbl{\f0 Aptos;}{\f1 Aptos Display;}}",
        (
            r"{\colortbl;"
            r"\red31\green78\blue121;"
            r"\red102\green102\blue102;"
            r"\red31\green78\blue121;"
            r"\red158\green173\blue186;}"
        ),
        r"\paperw12240\paperh15840\margl1080\margr1080\margt936\margb936",
        (
            r"\pard\qc\sa80\f1\fs36\b\cf1 "
            r"NeurIPS 2026 Reviewer Responses\b0\par"
        ),
        r"\pard\qc\sa240\f0\fs20\cf2 Submission 29735\par",
    ]

    for index, (reviewer, review, response) in enumerate(material):
        if index:
            content.append(r"\page")
        content.extend(
            [
                rf"\pard\sa60\f1\fs30\b\cf1 Reviewer {rtf_escape(reviewer)}\b0\par",
                r"\pard\brdrb\brdrs\brdrw10\brsp40\cf4\par",
                r"\pard\sb80\sa50\f0\fs17\b\cf2 ORIGINAL REVIEW\b0\par",
                rtf_paragraphs(
                    review,
                    font_size=17,
                    color=2,
                    before=0,
                    after=80,
                    line_spacing=220,
                ),
                r"\pard\sb80\brdrb\brdrs\brdrw10\brsp40\cf4\par",
                r"\pard\sb100\sa60\f0\fs24\b\cf3 AUTHOR RESPONSE\b0\par",
                rtf_paragraphs(
                    response,
                    font_size=23,
                    color=3,
                    before=0,
                    after=140,
                    line_spacing=280,
                ),
                r"\pard\sb80\brdrb\brdrs\brdrw10\brsp40\cf4\par",
            ]
        )

    content.append("}")
    DOC_PATH.write_text("\n".join(content), encoding="ascii")


def main() -> None:
    material = load_material()
    build_docx(material)
    build_doc(material)
    for path in (DOC_PATH, DOCX_PATH):
        print(f"{path.relative_to(DRAFT_DIR)}: {path.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
